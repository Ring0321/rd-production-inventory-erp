"""Load the patent ledger workbook into the isolated local ERP instance.

The source workbook is illustrative data, not evidence of factory operations.
This importer refuses a nonempty ERP database and the production DB port.
"""

from collections import Counter, defaultdict
from datetime import datetime
from decimal import Decimal
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile

import openpyxl
import pymysql


ROOT = Path(__file__).resolve().parents[2]
ENV_FILE = ROOT / "temp/local-runtime-patent/.env.local"
SOURCE_ZIP = Path(
    r"D:\芯珩科技\ERP\outputs\吨袋成品二维码批次流转仓储收发核验追溯系统_专利台账数据详细版_20260918_100002.zip"
)
MARKER = "专利演示数据（非现场实测）"
BUSINESS_DATE = "2026-09-18"


def insert(cur, table, values):
    columns = ", ".join(values)
    placeholders = ", ".join(["%s"] * len(values))
    cur.execute(
        f"INSERT INTO {table} ({columns}) VALUES ({placeholders})",
        tuple(values.values()),
    )
    return cur.lastrowid


def records(sheet):
    rows = sheet.values
    headers = next(rows)
    return [dict(zip(headers, row)) for row in rows]


def date_time(value):
    return datetime.fromisoformat(value) if isinstance(value, str) else value


def main():
    env = dict(
        line.split("=", 1)
        for line in ENV_FILE.read_text(encoding="utf-8").splitlines()
        if "=" in line and not line.startswith("#")
    )
    port = int(env["RD_DB_PORT"])
    if port != 3310 or "patent" not in str(ENV_FILE):
        raise RuntimeError("Refusing to import outside the isolated patent instance")
    with ZipFile(SOURCE_ZIP) as archive:
        xlsx = next(name for name in archive.namelist() if name.endswith(".xlsx"))
        book = openpyxl.load_workbook(BytesIO(archive.read(xlsx)), read_only=True, data_only=True)
    batches = records(book.worksheets[1])
    bags = records(book.worksheets[2])
    quality = records(book.worksheets[3])
    inbound = records(book.worksheets[4])
    outbound = records(book.worksheets[6])
    if len(batches) != 2 or len(bags) != 30 or len(inbound) != 30 or len(outbound) != 12:
        raise RuntimeError("Source workbook counts differ from the reviewed sample")

    db = pymysql.connect(
        host="127.0.0.1", port=port, user="root", password=env["RD_DB_PASSWORD"],
        database="rd_erp", charset="utf8mb4", autocommit=False,
    )
    try:
        with db.cursor() as cur:
            tables = (
                "rd_production_output", "rd_package_unit", "rd_stock_lot",
                "rd_stock_transaction", "rd_inbound_order", "rd_outbound_order",
            )
            for table in tables:
                cur.execute(f"SELECT COUNT(*) FROM {table}")
                if cur.fetchone()[0]:
                    raise RuntimeError(f"Refusing nonempty database: {table}")

            batch_ids = {}
            for batch in batches:
                batch_code = batch["批次号"]
                batch_ids[batch_code] = insert(cur, "rd_production_output", dict(
                    tenant_id=1, production_date=BUSINESS_DATE,
                    customer_name="专利演示客户", product_model=batch["产品编码"],
                    furnace_no=batch_code, furnace_layer=batch["生产线"],
                    output_qty=batch["实际赋码吨袋数"], package_weight=Decimal("1000"),
                    output_weight=Decimal(batch["实际赋码吨袋数"]) * 1000,
                    inbound_qty=batch["实际赋码吨袋数"], status="INBOUNDED",
                    source_sheet="01批次生产赋码",
                    remark=f"{MARKER}；源任务号 {batch['生产任务号']}；原始批次字段映射至炉号",
                    creator="1", updater="1",
                    create_time=date_time(batch["开始包装时间"]),
                    update_time=date_time(batch["结束包装时间"]),
                ))

            inbound_by_bag = {row["吨袋编号"]: row for row in inbound}
            outbound_by_bag = {row["吨袋编号"]: row for row in outbound}
            bag_by_code = {row["吨袋编号"]: row for row in bags}
            if len(inbound_by_bag) != 30 or len(outbound_by_bag) != 12:
                raise RuntimeError("Duplicate bag code in source movement sheets")

            location_codes = sorted({row["库位"] for row in inbound})
            location_ids = {}
            for code in location_codes:
                location_ids[code] = insert(cur, "rd_location", dict(
                    tenant_id=1, location_code=code, location_name=code,
                    warehouse_code="WH-01", warehouse_name="WH-01",
                    occupied_packages=sum(
                        1 for row in inbound if row["库位"] == code
                        and row["吨袋编号"] not in outbound_by_bag
                    ),
                    occupied_weight=1000 * sum(
                        1 for row in inbound if row["库位"] == code
                        and row["吨袋编号"] not in outbound_by_bag
                    ),
                    remark=MARKER, creator="1", updater="1",
                ))

            batch_counts = Counter(row["批次号"] for row in bags)
            outbound_counts = Counter(row["批次号"] for row in outbound)
            lot_ids = {}
            for batch in batches:
                code = batch["批次号"]
                current = batch_counts[code] - outbound_counts[code]
                lot_ids[code] = insert(cur, "rd_stock_lot", dict(
                    tenant_id=1, location_code="WH-01", customer_name="专利演示客户",
                    product_model=batch["产品编码"], furnace_no=code,
                    furnace_layer=batch["生产线"], inbound_qty=batch_counts[code],
                    outbound_qty=outbound_counts[code], warehouse_qty=current,
                    package_weight=1000, stock_weight=current * 1000,
                    last_biz_date=BUSINESS_DATE, status="AVAILABLE",
                    remark=MARKER, creator="1", updater="1",
                ))

            bag_ids = {}
            for bag in bags:
                code = bag["吨袋编号"]
                receipt = inbound_by_bag[code]
                location = receipt["库位"]
                bag_ids[code] = insert(cur, "rd_package_unit", dict(
                    tenant_id=1, package_code=code,
                    qr_code=bag["二维码内容摘要"],
                    production_output_id=batch_ids[bag["批次号"]],
                    production_date=BUSINESS_DATE,
                    customer_name="专利演示客户", product_name=bag["产品名称"],
                    product_model=bag["产品编码"], furnace_no=bag["批次号"],
                    furnace_layer=bag["生产线"], package_weight=Decimal("1000"),
                    location_id=location_ids[location], location_code=location,
                    stock_lot_id=lot_ids[bag["批次号"]],
                    status="OUTBOUNDED" if code in outbound_by_bag else "IN_STOCK",
                    version=1, remark=MARKER, creator="1", updater="1",
                    create_time=date_time(bag["包装完成时间"]),
                    update_time=date_time(
                        outbound_by_bag[code]["装车扫码时间"] if code in outbound_by_bag
                        else receipt["扫码入库时间"]
                    ),
                ))

            for check in quality:
                sampled = [code.strip() for code in check["抽检吨袋编号"].split(";")]
                for code in sampled:
                    insert(cur, "rd_quality_inspection", dict(
                        tenant_id=1,
                        inspection_no=f"{check['质检单号']}-{code[-3:]}",
                        package_id=bag_ids[code], status="PASSED", conclusion="PASS",
                        submitter_name=check["质检员"], reviewer_name=check["质检员"],
                        submit_time=date_time(check["放行/处置时间"]),
                        review_time=date_time(check["放行/处置时间"]),
                        version=1,
                        remark=f"{MARKER}；源批次抽检单 {check['质检单号']}；{check['检验结果']}；{check['批次处置']}",
                        creator="1", updater="1",
                        create_time=date_time(check["放行/处置时间"]),
                    ))

            inbound_orders = {}
            inbound_running = defaultdict(int)
            for receipt in inbound:
                order_no = receipt["入库单号"]
                if order_no not in inbound_orders:
                    rows = [row for row in inbound if row["入库单号"] == order_no]
                    inbound_orders[order_no] = insert(cur, "rd_inbound_order", dict(
                        tenant_id=1, order_no=order_no, biz_date=BUSINESS_DATE,
                        source_type="PRODUCTION_OUTPUT",
                        production_output_id=batch_ids[receipt["批次号"]],
                        total_packages=len(rows), total_weight=len(rows) * 1000,
                        status="POSTED", poster_name="专利演示导入",
                        post_time=max(date_time(row["扫码入库时间"]) for row in rows),
                        remark=MARKER, creator="1", updater="1",
                    ))
                code = receipt["吨袋编号"]
                bag = bag_by_code[code]
                batch = bag["批次号"]
                inbound_running[batch] += 1
                timestamp = date_time(receipt["扫码入库时间"])
                tx_id = insert(cur, "rd_stock_transaction", dict(
                    tenant_id=1, biz_date=BUSINESS_DATE, transaction_type="INBOUND",
                    stock_lot_id=lot_ids[batch], package_id=bag_ids[code],
                    location_id=location_ids[receipt["库位"]],
                    location_code=receipt["库位"], customer_name="专利演示客户",
                    product_model=bag["产品编码"], furnace_no=batch,
                    furnace_layer=bag["生产线"], quantity=1, package_weight=1000,
                    weight=1000, before_qty=inbound_running[batch] - 1,
                    after_qty=inbound_running[batch], source_type="INBOUND_ORDER",
                    source_no=order_no, source_line_no=code,
                    related_order_id=inbound_orders[order_no],
                    operator_name=receipt["入库操作人"],
                    idempotency_key=f"PATENT-IN-{code}",
                    remark=MARKER, creator="1", updater="1", create_time=timestamp,
                ))
                insert(cur, "rd_inbound_order_line", dict(
                    tenant_id=1, order_id=inbound_orders[order_no],
                    package_id=bag_ids[code], location_id=location_ids[receipt["库位"]],
                    location_code=receipt["库位"], quantity=1, weight=1000,
                    stock_lot_id=lot_ids[batch], stock_transaction_id=tx_id,
                    remark=MARKER, creator="1", updater="1", create_time=timestamp,
                ))

            outbound_orders = {}
            outbound_running = Counter()
            for shipment in outbound:
                number = shipment["发货单号"]
                code = shipment["吨袋编号"]
                bag = bag_by_code[code]
                batch = bag["批次号"]
                if number not in outbound_orders:
                    rows = [row for row in outbound if row["发货单号"] == number]
                    order_id = insert(cur, "rd_sale_order_extend", dict(
                        tenant_id=1, order_no=f"DEMO-SO-{number[-3:]}",
                        order_date=BUSINESS_DATE,
                        customer_name=shipment["客户代号"],
                        product_model=bag["产品编码"], required_qty=len(rows),
                        allocated_qty=len(rows), fulfilled_qty=len(rows),
                        status="OPEN", remark=f"{MARKER}；对应发货单 {number}",
                        creator="1", updater="1",
                    ))
                    fulfillment_id = insert(cur, "rd_order_fulfillment", dict(
                        tenant_id=1, order_id=order_id, stock_lot_id=lot_ids[batch],
                        matched_qty=len(rows), allocated_qty=len(rows),
                        outbound_qty=len(rows), source_type="DEMO",
                        status="OUTBOUNDED", remark=MARKER, creator="1", updater="1",
                    ))
                    outbound_id = insert(cur, "rd_outbound_order", dict(
                        tenant_id=1, outbound_no=number,
                        outbound_date=BUSINESS_DATE, order_id=order_id,
                        fulfillment_id=fulfillment_id, stock_lot_id=lot_ids[batch],
                        quantity=len(rows), package_weight=1000, weight=len(rows) * 1000,
                        status="POSTED", version=1,
                        poster_name="专利演示导入",
                        post_time=max(date_time(row["装车扫码时间"]) for row in rows),
                        remark=MARKER, creator="1", updater="1",
                    ))
                    outbound_orders[number] = outbound_id
                outbound_running[batch] += 1
                location = inbound_by_bag[code]["库位"]
                timestamp = date_time(shipment["装车扫码时间"])
                tx_id = insert(cur, "rd_stock_transaction", dict(
                    tenant_id=1, biz_date=BUSINESS_DATE, transaction_type="OUTBOUND",
                    stock_lot_id=lot_ids[batch], package_id=bag_ids[code],
                    location_id=location_ids[location], location_code=location,
                    customer_name="专利演示客户", product_model=bag["产品编码"],
                    furnace_no=batch, furnace_layer=bag["生产线"],
                    quantity=1, package_weight=1000, weight=1000,
                    before_qty=batch_counts[batch] - outbound_running[batch] + 1,
                    after_qty=batch_counts[batch] - outbound_running[batch],
                    source_type="OUTBOUND_ORDER", source_no=number,
                    source_line_no=code, related_order_id=outbound_orders[number],
                    operator_name="专利演示导入", idempotency_key=f"PATENT-OUT-{code}",
                    remark=MARKER, creator="1", updater="1", create_time=timestamp,
                ))
                insert(cur, "rd_outbound_order_line", dict(
                    tenant_id=1, order_id=outbound_orders[number],
                    package_id=bag_ids[code], package_code=code,
                    stock_lot_id=lot_ids[batch], stock_transaction_id=tx_id,
                    location_id=location_ids[location], location_code=location,
                    customer_name="专利演示客户", product_model=bag["产品编码"],
                    furnace_no=batch, furnace_layer=bag["生产线"],
                    package_weight=1000, scan_device="PDA-01",
                    scanner_name="专利演示导入", creator="1", updater="1",
                    create_time=timestamp,
                ))

        db.commit()
        print("Imported 2 batches, 30 bags, 8 sampled QC records, 30 inbound and 12 outbound movements")
    except Exception:
        db.rollback()
        raise
    finally:
        db.close()


if __name__ == "__main__":
    main()
